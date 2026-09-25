"""docx 解析 + 结构化切分：
章节层级靠正文编号前缀表达（"1　文档说明"、"6.1　GMV"）
表格是口径/字段定义的主要载体，必须整体保留为一个 chunk 避免破坏语义
"""

import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

HEADING_RE = re.compile(r"^(\d+(?:\.\d+)*)\s+(.+)$")

CHUNK_TEXT = "text"     # 正文段落chunk
CHUNK_TABLE = "table"   # 表格chunk

@dataclass
class Chunk:
    """知识库最小单元，section_path是命中后溯源到文档章节的路径"""
    doc_title: str  # 整个文档的标题
    section_path: list[str]     # 章节的路径栈
    content: str    # 正文内容
    chunk_type: str = CHUNK_TEXT  # text/table

def _table_to_md(table: Table) -> str:
    """整表转markdown：第一行为表头其余为数据行
    防止python-docx把合并位置的cell复用同一个对象，文本内容重复，用“相邻重复去重”解决
    """
    rows = []
    for row in table.rows:
        cells = [c.text.strip() for c in row.cells]
        dedup = [v for i, v in enumerate(cells) if i == 0 or v != cells[i - 1]]     # 保留第0列和相邻不重复的列
        rows.append(dedup)
    width = max(len(r) for r in rows)
    sep = ["---"] * width
    lines = ["| " + " | ".join(rows[0]) + " |",
             "| " + " | ".join(sep) + " |"]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)

def parse_docx(path: str | Path) -> list[Chunk]:
    """把docx切成一串Chunk:正文按编号章节聚合，表格单独成chunk"""
    doc = Document(str(path))
    doc_title = doc.paragraphs[0].text.strip() or Path(path).stem
    chunks: list[Chunk] = []
    section: list[str] = []  # 当前标题栈，如 ["6　核心指标口径定义", "6.1　GMV（商品交易总额）"]
    buf: list[str] = [] # 文本缓冲区  

    def flush():
        # 文档大标题直接丢弃(已经存入doc_title)
        if buf and section:
            chunks.append(Chunk(doc_title, list(section), "\n".join(buf)))
        buf.clear()

    # doc.paragraphs 和 doc.tables 不能分开遍历，否则表格会挂错章节
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):  # 是段落
            text = Paragraph(child, doc).text.strip()   
            m = HEADING_RE.match(text)  # 匹配编号标题如 6　核心指标口径定义
            if m:
                flush()
                depth = m.group(1).count(".") + 1   # group(1)捕获数字，group(2)是文字部分，这行是算出标题是几级的
                section = section[:depth - 1] + [text]  # 同级新标题替换旧标题，截断掉原来低级的
            elif text:  # 普通正文
                buf.append(text)
        elif child.tag == qn("w:tbl"):  # 是表格
            flush()
            chunks.append(Chunk(doc_title, list(section), _table_to_md(Table(child, doc)), chunk_type=CHUNK_TABLE))
    flush()
    return chunks

if __name__ == "__main__":
    here = Path(__file__).parent
    for path in sorted(here.glob("*.docx")):
        chunks = parse_docx(path)
        print(f"\n=== {path.name}:{len(chunks)}个chunk ===")
        for c in chunks: 
            loc = "/".join(c.section_path) or "(无章节)"
            print(f"    [{c.chunk_type}] {loc} | {len(c.content)}字")
        text_c = next((c for c in chunks if c.chunk_type == CHUNK_TEXT), None)
        table_c = next((c for c in chunks if c.chunk_type == CHUNK_TABLE), None)
        if text_c:
            print(f"\n  --- 抽查 text chunk（{text_c.section_path[-1]}）---")
            print("  " + text_c.content[:200])
        if table_c:
            print(f"\n  --- 抽查 table chunk（{table_c.section_path[-1]}）---")
            print("\n".join("  " + line for line in table_c.content.splitlines()[:6]))